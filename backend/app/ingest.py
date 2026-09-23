"""Turn uploaded files into plain text. Supports pdf, docx, txt, md, csv, json."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="ignore"))))
        return "\n".join(" | ".join(r) for r in rows)
    if ext == ".json":
        return json.dumps(json.loads(data.decode("utf-8", errors="ignore")), ensure_ascii=False, indent=1)
    return data.decode("utf-8", errors="ignore")


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    """Paragraph-aware chunking by characters."""
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) + 1 <= size:
            buf = f"{buf}\n{p}" if buf else p
            continue
        if buf:
            chunks.append(buf)
        while len(p) > size:
            chunks.append(p[:size])
            p = p[size - overlap :]
        buf = p
    if buf:
        chunks.append(buf)
    return chunks
